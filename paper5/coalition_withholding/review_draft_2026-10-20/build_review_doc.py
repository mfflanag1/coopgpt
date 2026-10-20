from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
ROOT=Path(__file__).parent
OUT=ROOT/'Coalition_withholding_specialist_review_draft.docx'
doc=Document()
s=doc.sections[0]
s.top_margin=s.bottom_margin=Inches(.75)
s.left_margin=s.right_margin=Inches(.82)
s.page_width=Inches(8.5);s.page_height=Inches(11)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','List Bullet','List Number']:
    st=doc.styles[name];st.font.name='Cambria';st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(11 if name not in ['Title','Heading 1','Heading 2'] else {'Title':23,'Heading 1':16,'Heading 2':12}[name])
    st.paragraph_format.space_after=Pt(7)
    st.paragraph_format.line_spacing=1.08
for name in ['Title','Heading 1','Heading 2']:
    doc.styles[name].paragraph_format.keep_with_next=True
    doc.styles[name].paragraph_format.space_before=Pt(8)
doc.core_properties.author='Maxwell Flanagan'
doc.core_properties.title='Coalition withholding specialist review packet'
doc.core_properties.subject='Correctness and priority review request'
doc.core_properties.keywords='coalition withholding, production, cost sharing, review'
def p(text='',style=None):
    import re
    para=doc.add_paragraph(style=style)
    last=0
    for m in re.finditer(r'([_^])(?:\(([^)]*)\)|([A-Za-z0-9]+))',text):
        para.add_run(text[last:m.start()])
        run=para.add_run(m.group(2) or m.group(3))
        if m.group(1)=='_':run.font.subscript=True
        else:run.font.superscript=True
        last=m.end()
    para.add_run(text[last:])
    return para
def h(text):return doc.add_heading(text,level=1)
def h2(text):return doc.add_heading(text,level=2)
def link(label,url):
    para=doc.add_paragraph();a=OxmlElement('w:hyperlink');a.set(qn('r:id'),para.part.relate_to(url,RT.HYPERLINK,is_external=True))
    r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'174A74');pr.append(color);r.append(pr)
    t=OxmlElement('w:t');t.text=label;r.append(t);a.append(r);para._p.append(a)
    return para
def eq(text):
    import re
    def mr(value):
        r=OxmlElement('m:r');t=OxmlElement('m:t');t.text=value;r.append(t);return r
    para=doc.add_paragraph();para.alignment=WD_ALIGN_PARAGRAPH.CENTER
    om=OxmlElement('m:oMath');nodes=[];i=0
    while i<len(text):
        if text[i] in '_^' and nodes:
            match=re.match(r'([_^])(?:\(([^)]*)\)|([A-Za-z0-9]+))',text[i:])
            if match:
                typ='Sub' if match.group(1)=='_' else 'Sup'
                node=OxmlElement('m:s'+typ);e=OxmlElement('m:e');e.append(nodes.pop());node.append(e)
                child=OxmlElement('m:'+typ.lower());child.append(mr(match.group(2) or match.group(3)));node.append(child);nodes.append(node)
                i+=len(match.group());continue
        nodes.append(mr(text[i]));i+=1
    for node in nodes:om.append(node)
    para._p.append(om)
    return para
base='https://github.com/mfflanag1/coopgpt/blob/2db8290fd7d213ec2265e8866eb68a5d3b99bd96/paper5/coalition_withholding/'
p('Coalition withholding specialist review packet','Title')
p('Maxwell Flanagan  |  20 October 2026','Subtitle')
h('Cover note for Marina Núñez')
p('Subject: Request for correctness and priority review of coalition withholding results')
p('Dear Professor Núñez,')
p('I would appreciate your assessment of my manuscript, “The cost of coalition withholding: resource dimension and coalition size.” It studies how much external funding is needed to cover every coalition’s production while preventing coordinated withholding of resources by fixed owners.')
p('The main new result in this research program is an unrestricted lower bound of order √s for technologies in which each activity uses at most three resources, where s is the maximum protected coalition size. This improves our earlier lower exponent from 1/4 to 1/2. The general upper bound remains 3s². The paper also consolidates sharp resource-dimension bounds and includes complete proofs, a higher-arity extension, and exact reproducibility checks.')
p('I am seeking a focused correctness and priority review. In particular, I would value scrutiny of the fixed-owner symmetrization and payment-ladder argument, and advice on whether existing strong group monotonicity, absence-proofness, or cost-sharing results already imply these quantitative claims. The accompanying summary identifies the contract and the most consequential questions.')
p('The work has used AI-assisted research and internal auditing. Those checks have not established historical novelty or replaced human specialist review. I would welcome a specific objection, an antecedent theorem with a valid transfer to this setting, or guidance on whether the results merit a single paper. If a complementary cost-sharing or algorithmic-game-theory reader would be better placed to assess part of the work, a referral would also be helpful.')
p('Thank you for considering it.\n\nMaxwell Flanagan')
h2('Files to accompany the request')
p('Attach the 26-page manuscript PDF and the three-page specialist brief. The source matrix and complete replay package can be supplied alongside them. The repository is private, so attachments are preferable unless the recipient already has access.')
link('Manuscript at the reviewed commit',base+'manuscript.pdf')
link('Specialist brief at the reviewed commit',base+'specialist_brief.pdf')
doc.add_page_break()
h('The exact mathematical contract')
p('A fixed institution specifies finitely many actual owners, resource coordinates, a public production function F, and coordinatewise endowment boxes. Each owner reports any real bundle bᵢ between zero and its endowment. The owner roster and eligibility remain fixed throughout a deviation. Write b_C for the sum of the reports of coalition C, and y_C for its total payment.')
p('At every report profile, payments are nonnegative and satisfy both conditions:')
eq('y_C(b) ≥ F(b_C);     ∑_i y_i(b) ≤ κ F(b_N).')
p('For a coalition C with at most s members, let δ denote a coordinatewise feasible withdrawal supported on C. Unchanged members may remain inside the protected coalition. The two incentive contracts are:')
eq('TCWM:  y_C(b) − y_C(b − δ) ≥ 0.')
eq('Full OUT:  y_C(b) − y_C(b − δ) ≥ F(δ_C).')
p('TCWM protects the coalition’s aggregate payment. Full OUT also protects production from the jointly withdrawn resources. The factor κ bounds total funding; the corresponding external supplement is at most (κ − 1)F(b_N). Coverage applies to every coalition, even when only coalitions of size at most s receive incentive protection.')
p('The lower bounds allow arbitrary admissible payment rules. They do not assume anonymity, continuity, linearity, zero payments to inactive owners, a marginal-payment floor, or representation by efficient production prices.')
h('Claims for review')
h2('Sharp dependence on resource dimension')
p('For finite-valued monotone concave degree-one production on d resource types, protecting all coalitions gives the following worst-instance total-funding factors:')
eq('D_d = (d + 1)/2;     D_d^OUT = 1 + c(d − 1),     c(2 − exp(−1/c)) = 1.')
p('These are the consolidated dimension results in Section 3.')
h2('Dependence on activity arity and coalition size')
p('For finite nonnegative packing production, let r be the maximum number of positive resource entries in a profitable activity. Owner count and resource dimension may both grow. Section 4 proves the full-OUT upper r s^(r−1). Section 5 supplies the new unrestricted arity-three lower:')
eq('√(s/6) ≤ K₃(s) ≤ K₃^OUT(s) ≤ 3s².')
p('The sharper certificate B_s is asymptotic to √(2πs)/(3√e), with B₆ = 493/243. For fixed r ≥ 3, Appendix A gives the higher-arity lower exponent (r − 2)/(r − 1) and a separate pure-owner realization. The gap remains open. The witnesses have exponentially many resource rows; the pure-owner construction also needs more owners than the fully eligible construction.')
doc.add_page_break()
h('Proof structure and correctness questions')
p('The arity-three construction uses nonzero binary vectors as resource rows and unit-profit activities {a, b, a + b}. A finite family of profiles has production value one. A legal finite average over owner/resource permutations reduces arbitrary payment rules to a symmetric comparison without changing the fixed institution or the real report boxes.')
p('Coverage forces a positive payment at a seed profile. Restoring one low owner while protecting it together with selected unchanged private owners produces a recurrence. Its product forces aggregate payments to grow at square-root order. An asymmetric original rule must be expensive somewhere in the finite orbit; the displayed representative need not itself be expensive.')
p('For a capped coalition size, the recurrence uses k_m = min{s − 1, 2m − 2} and the following inequality:')
eq('(k_m + 1)q_m ≥ k_m q_(m−1).')
p('The old restorer’s payment is discarded using nonnegativity, never by assuming it is zero. Optimizing the resulting certificate gives B_s; optimality of that ladder length is not optimality among all possible lower-bound arguments.')
h2('The highest consequence correctness checks')
for text in [
'Check that every averaging operation preserves the full report domain and eligibility, including dormant high-resource reports, zero-report owners, and the joint resource/owner action in the pure-owner version.',
'Check the seed stabilizers, private-owner transitivity, and the claim that the listed activities are the only supported columns of the full production LP at each witness profile.',
'Check coalition-size accounting and every retained payment term. In particular, no individual-incumbent monotonicity or dummy-zero condition should enter the arbitrary-rule lower.',
'For the dimension lower, check finite-owner compactness, dilation averaging, inactive payments and shifted abundant totals without assuming regularity of the original payment rules.',
'For the upper bounds, check activation at support boundaries, the sparse coefficient signs, and simultaneous validity of all-coalition coverage and incentive protection on the same report domain.'
]:p(text,'List Bullet')
h2('What a useful review would return')
p('Please identify any fatal gap or missing hypothesis, distinguish repairable issues from invalid claims, and point to exact antecedent statements where possible. An assessment of whether the accepted results form a worthwhile single contribution would be valuable. This request does not assume that the paper resolves a recognized major open problem.')
doc.add_page_break()
h('Priority comparisons and reproducibility')
p('Strong group monotonicity, absence-proofness, projective constructions, static cost of stability, bounded-rank positive hypergraphs and binomial marginal values have established antecedents. The proposed contribution is the quantitative funding frontier across reports under the stated contract, despite efficient static coalition covers at each packing profile.')
h2('The closest comparisons')
p('Moulin and Sprumont, On Demand Responsiveness in Additive Cost Sharing: the inspected 2004 author version has the relevant aggregate-sum monotonicity axiom. Its Proposition 2(ii) also imposes additivity, separability and exact balance. Please assess whether a theorem or valid reduction nevertheless transfers to the present overfunded production setting.')
p('Emre Doğan, Absence-proofness: the inspected 2013 working paper has a close binary-participation analogue of full OUT with efficient subgame allocations. Comparison with the published article and any quantitative subsidized extension remains open.')
p('Immorlica, Mahdian and Mirrokni, Limitations of cross-monotonic cost-sharing schemes: the profit-sharing lower uses per-incumbent cross-monotonicity. That stronger condition must not be silently substituted for the aggregate coalition-sum inequality used here.')
p('Published-source comparisons remain incomplete for Sprumont’s nearly-serial sharing, serial-method axiomatization and discrete Aumann–Shapley/multigood work. Schoppmann’s small-coalition cost sharing and Moulin–Sprumont’s cross-subsidization work remain additional leads. The detailed matrix records inspected versions and unresolved questions; missing sources are not evidence of novelty.')
h2('Specific priority questions')
for text in [
'Does an existing theorem imply either sharp dimension endpoint or the sparse coalition-size lower while preserving the owners, real reports, coverage, incentive direction and total-funding normalization?',
'Is there already an unrestricted lower with exponent greater than 1/2 at fixed activity arity under this exact contract?',
'Do the published originals change any comparison made from an earlier author or working version? Which incremental theorem, if any, remains after those comparisons?'
]:p(text,'List Bullet')
h2('Version and exact checks')
p('This packet refers to the 18 October manuscript consolidated at Git commit 2db8290fd7d213ec2265e8866eb68a5d3b99bd96, merged into main on 20 October 2026. The complete package contains standalone LaTeX, proofs, review dispositions, source comparisons, SHA256SUMS and six focused Python standard-library replay suites.')
p('Run python3 reproduce.py from the extracted package with assertions enabled. The reproduction notes specify the finite ranges and exclusions. The recorded fresh-copy replays passed, but finite certificates do not establish the all-size theorem or historical priority. Internal model reviews are not human peer review.')
link('Detailed priority matrix',base+'review/PRIORITY_MATRIX.txt')
link('Reproduction instructions and scope',base+'README.txt')
link('Complete review ZIP','https://github.com/mfflanag1/coopgpt/blob/2db8290fd7d213ec2265e8866eb68a5d3b99bd96/paper5/coalition_withholding-review-2026-10-18.zip')
for paragraph in doc.paragraphs:
    paragraph.paragraph_format.widow_control=True
footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Coalition withholding review packet  •  ');r.font.size=Pt(9)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');footer._p.append(f)
for style in doc.styles:
    for tag in ('w:pBdr',):
        for element in style.element.findall('.//'+qn(tag)):
            element.getparent().remove(element)
    for fonts in style.element.findall('.//'+qn('w:rFonts')):
        for attr in list(fonts.attrib):
            if 'theme' in attr.lower():del fonts.attrib[attr]
for el in doc.element.findall('.//'+qn('w:pBdr')):
    el.getparent().remove(el)
doc.save(OUT)
print(OUT)
