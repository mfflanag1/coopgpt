"""Read-only availability inventory; this is not a forecast backtest."""
from pathlib import Path
from collections import Counter
import hashlib
import json
import re
import openpyxl

p = Path(__file__).with_name("2026-accuracy-report.xlsx")
expected = "e0f24fe03a79af016826089c653b2c2e15e55edcb04a7bd4be5fa643136425fc"
if hashlib.sha256(p.read_bytes()).hexdigest() != expected:
    raise ValueError("Source changed: inspect the new vintage before using this inventory")
w = openpyxl.load_workbook(p, read_only=True, data_only=True)
records = []
hist_sheets = []
for sheet in w:
    rows = list(sheet.iter_rows(values_only=True))
    if tuple(rows[0]) != ("Zone", "Year", "Metric", "Load MW"):
        raise ValueError(f"Unexpected headers: {sheet.title}")
    history = {int(year): (mw, i) for i, (_, year, metric, mw) in enumerate(rows[1:], 2)
               if metric == "Historical July Peak" and isinstance(year, (int, float))
               and isinstance(mw, (int, float))}
    if history:
        hist_sheets.append(sheet.title)
    for i, (_, year, metric, mw) in enumerate(rows[1:], 2):
        match = re.fullmatch(r"(\d{4}) Load Forecast July Peak", str(metric))
        if match and isinstance(year, (int, float)) and isinstance(mw, (int, float)):
            vintage = int(match[1])
            if year in history and year >= vintage:
                records.append({"sheet": sheet.title, "vintage": vintage,
                                "target": int(year), "horizon": int(year) - vintage,
                                "forecast_cell": f"D{i}",
                                "historical_cell": f"D{history[year][1]}"})
zones = {"AEP", "COMED", "DOM", "PS"}
zone_records = [r for r in records if r["sheet"] in zones]
print(json.dumps({
    "warning": "Label matching only. Original vintage fidelity, stable populations and revised outcomes remain unchecked. Aggregates overlap; positive-horizon pairs are not independent observations.",
    "sha256": expected, "sheets": len(w.sheetnames),
    "sheets_with_history": hist_sheets,
    "nonoverlapping_zone_sheets": sorted(zones),
    "zone_pairs_by_horizon": dict(sorted(Counter(r["horizon"] for r in zone_records).items())),
    "zone_pair_cells": zone_records,
}, indent=2))
